"""
Model training pipeline for the SEED EEG emotion recognition task.

This module mirrors the original `3model.py` script, reorganised into a package
structure without changing the underlying algorithmic steps. It loads the
preprocessed EEG chunks, reprograms them via a BERT backbone, and performs
cross-validation alongside a held-out test split.
"""

import copy
import json
import hashlib
import math
import os
import argparse
import time
from collections import Counter
from pathlib import Path

import numpy as np
import torch
import torch.nn as nn
from openpyxl import Workbook, load_workbook
from sklearn.metrics import (accuracy_score, classification_report,
                             confusion_matrix, f1_score, precision_score,
                             recall_score)
from sklearn.model_selection import KFold, train_test_split
from torch.optim import AdamW
from torch.utils.data import DataLoader, Dataset, TensorDataset
from transformers import BertModel, BertTokenizer


DEFAULT_DATA_DIR = "/home/aispeech/codes/zxy/SEED_chunks"
DEFAULT_BERT_DIR = "/home/aispeech/codes/zxy/TPANet-main/TPANet-main2/models/bert-base-uncased"


def _is_git_lfs_pointer(file_path: Path) -> bool:
    if not file_path.is_file():
        return False
    try:
        if file_path.stat().st_size > 2048:
            return False
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return False
    return text.startswith("version https://git-lfs.github.com/spec/v1")


def _validate_local_bert_weights(bert_path: str) -> None:
    model_dir = Path(bert_path)
    if not model_dir.is_dir():
        return
    pytorch_bin = model_dir / "pytorch_model.bin"
    safetensors_file = model_dir / "model.safetensors"
    if _is_git_lfs_pointer(pytorch_bin):
        if safetensors_file.is_file() and safetensors_file.stat().st_size > 1024 * 1024:
            return
        raise RuntimeError(
            f"检测到无效模型权重: {pytorch_bin}\n"
            "该文件是 Git LFS 指针，不是真实的 BERT 权重。\n"
            "请下载完整权重到该目录（任选其一）：\n"
            "  - pytorch_model.bin\n"
            "  - model.safetensors\n"
            "下载完成后重新运行脚本即可。"
        )


def _resolve_bert_source(bert_arg: str) -> str:
    raw_value = str(bert_arg).strip()
    if not raw_value:
        raise ValueError("参数 --bert 不能为空。")

    expanded_path = Path(raw_value).expanduser()
    if not expanded_path.is_dir():
        raise FileNotFoundError(
            f"未找到 BERT 本地目录: {expanded_path}\n"
            "请确认路径存在，或通过 --bert 显式指定有效本地目录。"
        )
    return str(expanded_path)


def _is_cuda_oom_error(exc: Exception) -> bool:
    return isinstance(exc, RuntimeError) and "out of memory" in str(exc).lower()


def _sync_cuda_if_needed(device) -> None:
    if device.type == "cuda":
        torch.cuda.synchronize()


def _append_jsonl(file_path: Path, record: dict) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    with open(file_path, "a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")


class XlsxTextLogger:
    def __init__(self, file_path):
        self.file_path = Path(file_path)
        self.sheet_name = "logs"
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        if self.file_path.exists():
            self.workbook = load_workbook(self.file_path)
        else:
            self.workbook = Workbook()

        if self.sheet_name in self.workbook.sheetnames:
            self.worksheet = self.workbook[self.sheet_name]
        else:
            self.worksheet = self.workbook.active
            self.worksheet.title = self.sheet_name
            self.worksheet.append(["timestamp", "message"])
            self._flush()

    def write(self, text):
        if text is None:
            return
        text = str(text).replace("\r\n", "\n")
        if not text:
            return
        lines = text.split("\n")
        for line in lines:
            timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
            self.worksheet.append([timestamp, line])
        self._flush()

    def _flush(self):
        self.workbook.save(self.file_path)

    def close(self):
        self._flush()
        self.workbook.close()

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()


class EEGDataset(Dataset):
    def __init__(self, file_path):
        self.file_path = file_path
        with np.load(file_path, allow_pickle=True) as npz_file:
            self.data = npz_file['data']
            labels = npz_file['labels']
            labels = np.array(labels, dtype=np.int64)
            if labels.size == 0:
                raise ValueError(f"{file_path} 中 labels 为空。")
            labels = labels - labels.min()
            self.labels = labels

    def __len__(self):
        return len(self.labels)

    def __getitem__(self, idx):
        data_chunk = self.data[idx]
        label = self.labels[idx]
        return torch.tensor(data_chunk, dtype=torch.float32), torch.tensor(label, dtype=torch.int64)


class PositionalEmbedding(nn.Module):
    def __init__(self, d_model, max_len=5000):
        super(PositionalEmbedding, self).__init__()
        pe = torch.zeros(max_len, d_model).float()
        position = torch.arange(0, max_len).float().unsqueeze(1)
        div_term = (torch.arange(0, d_model, 2).float() * -(math.log(10000.0) / d_model)).exp()
        pe[:, 0::2] = torch.sin(position * div_term)
        pe[:, 1::2] = torch.cos(position * div_term)
        pe = pe.unsqueeze(0)
        self.register_buffer('pe', pe)

    def forward(self, x):
        return self.pe[:, :x.size(1), :].to(x.device)


class PatchEmbedding(nn.Module):
    def __init__(self, patch_len, d_model, stride, num_channels, dropout=0.1):
        super(PatchEmbedding, self).__init__()
        self.patch_len = patch_len
        self.stride = stride
        self.dropout = nn.Dropout(dropout)
        self.num_channels = num_channels
        self.value_embedding = nn.Conv1d(in_channels=num_channels, out_channels=d_model, kernel_size=1)

    def forward(self, x):
        x = x.unfold(dimension=-1, size=self.patch_len, step=self.stride)
        x = x.permute(0, 2, 1, 3).contiguous()
        x = x.view(-1, self.num_channels, self.patch_len)
        x = self.value_embedding(x)
        x = x.permute(0, 2, 1)
        return self.dropout(x)


class ReprogrammingLayer(nn.Module):
    def __init__(self, embed_dim, llm_embed_dim, num_heads=8, max_len=5000):
        super(ReprogrammingLayer, self).__init__()
        self.linear = nn.Linear(embed_dim, llm_embed_dim)
        self.ln = nn.LayerNorm(llm_embed_dim)
        self.multihead_attn = nn.MultiheadAttention(embed_dim=llm_embed_dim, num_heads=num_heads, batch_first=True)
        self.positional_embedding = PositionalEmbedding(d_model=llm_embed_dim, max_len=max_len)

    def forward(self, target_embedding, source_embedding, value_embedding):
        batch_size = source_embedding.size(0)
        num_patches = target_embedding.size(0) // batch_size

        target_embedding = target_embedding.view(
            batch_size,
            num_patches,
            target_embedding.size(1),
            target_embedding.size(2)
        )

        target_embedding = target_embedding.mean(dim=2)
        target_embedding = self.linear(target_embedding)

        positional_encoding = self.positional_embedding(target_embedding)
        target_embedding = target_embedding + positional_encoding

        target_embedding = self.ln(target_embedding)

        attn_output, attn_weights = self.multihead_attn(
            target_embedding, source_embedding, value_embedding,
            need_weights=True,
            average_attn_weights=False,
        )

        return attn_output, attn_weights


class ClassificationHead(nn.Module):
    def __init__(self, llm_embed_dim, num_labels):
        super(ClassificationHead, self).__init__()
        self.dropout = nn.Dropout(0.5)
        self.classifier = nn.Linear(llm_embed_dim, num_labels)

    def forward(self, x):
        x = self.dropout(x)
        logits = self.classifier(x)
        return logits


def generate_statistics(sample_trial):
    sample_trial_flat = sample_trial.view(sample_trial.size(0), -1)
    min_values = sample_trial_flat.min(dim=1)[0]
    max_values = sample_trial_flat.max(dim=1)[0]
    median_values = sample_trial_flat.median(dim=1)[0]
    trends = sample_trial_flat[:, -1] - sample_trial_flat[:, 0]
    return min_values, max_values, median_values, trends


def generate_prompts(batch_size, min_values, max_values, median_values, trends, mode="original", rng=None):
    """
    Build natural-language prompts from simple EEG statistics.

    Modes are used for reviewer-required ablations:
      - original: use true (min/max/median/trend)
      - generic: remove statistics (same capacity, less semantic guidance)
      - shuffle: keep numbers but shuffle their roles (break semantics, keep format)
      - random: keep format but replace numbers with random values in a plausible range
    """
    if rng is None:
        rng = np.random.RandomState(42)

    prompts = []
    for i in range(batch_size):
        min_value = float(min_values[i].item())
        max_value = float(max_values[i].item())
        median_value = float(median_values[i].item())
        trend_value = float(trends[i].item())
        trend = 'upward' if trend_value > 0 else 'downward'

        if mode == "generic":
            prompt = (
                "Based on the provided EEG data, classify the underlying emotion. "
                "The EEG signal shows certain temporal dynamics and amplitude variations."
            )
        elif mode == "shuffle":
            stats = [min_value, max_value, median_value]
            rng.shuffle(stats)
            min_s, max_s, med_s = stats
            prompt = (
                "Based on the provided EEG data, classify the underlying emotion. "
                f"Statistics: minimum value = {min_s:.2f}, maximum value = {max_s:.2f}, "
                f"median value = {med_s:.2f}. The overall trend of the data is {trend}."
            )
        elif mode == "random":
            lo, hi = (min_value, max_value) if min_value < max_value else (max_value, min_value + 1e-3)
            min_s = rng.uniform(lo, hi)
            max_s = rng.uniform(lo, hi)
            med_s = rng.uniform(lo, hi)
            trend = rng.choice(['upward', 'downward'])
            prompt = (
                "Based on the provided EEG data, classify the underlying emotion. "
                f"Statistics: minimum value = {min_s:.2f}, maximum value = {max_s:.2f}, "
                f"median value = {med_s:.2f}. The overall trend of the data is {trend}."
            )
        else:  # original
            prompt = (
                "Based on the provided EEG data, classify the underlying emotion. "
                f"Statistics: minimum value = {min_value:.2f}, maximum value = {max_value:.2f}, "
                f"median value = {median_value:.2f}. The overall trend of the data is {trend}."
            )

        prompts.append(prompt)
    return prompts




def save_attention_heatmap(attn_weights, prompts, save_path, token_limit=30):
    """Save a simple cross-attention heatmap (avg over heads) for the first sample in a batch.

    attn_weights: Tensor (B, num_heads, P, L)
    """
    import matplotlib.pyplot as plt

    w = attn_weights[0].mean(dim=0).detach().cpu().numpy()  # (P, L)
    w = w[:, :token_limit]

    plt.figure()
    plt.imshow(w, aspect='auto')
    plt.colorbar()
    plt.xlabel("Prompt token index")
    plt.ylabel("EEG patch index")
    plt.title("Cross-attention (avg over heads)")
    plt.tight_layout()
    plt.savefig(save_path, dpi=200)
    plt.close()


def save_attention_records(attn_blocks, token_id_blocks, label_blocks, save_path):
    """Persist aggregated per-sample attention tensors for post-hoc interpretability plots."""
    if save_path is None or (not attn_blocks):
        return

    out_path = Path(save_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    attn = np.concatenate(attn_blocks, axis=0).astype(np.float16)
    token_ids = np.concatenate(token_id_blocks, axis=0).astype(np.int32)
    labels = np.concatenate(label_blocks, axis=0).astype(np.int64)
    np.savez_compressed(out_path, attn=attn, token_ids=token_ids, labels=labels)


def profile_efficiency(model_components, device, batch_size=32, iters=200, warmup=50, prompt_mode="original", soft_prompt=None, prompt_cache_cpu=None, use_cached_prompts=False):
    """Measure params + inference latency (per forward) + peak GPU memory."""
    import time

    bert_model, patch_embedding, reprogramming_layer, classification_head = model_components
    bert_model.eval(); patch_embedding.eval(); reprogramming_layer.eval(); classification_head.eval()

    total_params = 0
    trainable_params = 0
    for m in [bert_model, patch_embedding, reprogramming_layer, classification_head]:
        for p in m.parameters():
            n = p.numel()
            total_params += n
            if p.requires_grad:
                trainable_params += n

    if device.type == "cuda":
        import torch
        torch.cuda.reset_peak_memory_stats()

    import torch
    x = torch.randn(batch_size, 62, 1000, device=device)
    min_values, max_values, median_values, trends = generate_statistics(x)
    prompts = generate_prompts(batch_size, min_values, max_values, median_values, trends, mode=prompt_mode)

    if prompt_mode == "soft":
        if soft_prompt is None:
            raise ValueError("soft_prompt must be provided when prompt_mode='soft'")
        prompt_embeddings = soft_prompt.expand(batch_size, -1, -1)
    else:
        if use_cached_prompts and (prompt_cache_cpu is not None):
            prompt_embeddings = prompt_cache_cpu[:batch_size].to(device).float()
        else:
            prompt_inputs = tokenizer(
                prompts, return_tensors="pt", padding="max_length", truncation=True, max_length=50
            ).to(device)
            with torch.no_grad():
                prompt_embeddings = bert_model(**prompt_inputs).last_hidden_state

    with torch.no_grad():
        for _ in range(warmup):
            eeg_emb = patch_embedding(x)
            out, _ = reprogramming_layer(eeg_emb, prompt_embeddings, prompt_embeddings)
            pooled = out.mean(dim=1)
            _ = classification_head(pooled)

        if device.type == "cuda":
            torch.cuda.synchronize()
        start = time.time()
        for _ in range(iters):
            eeg_emb = patch_embedding(x)
            out, _ = reprogramming_layer(eeg_emb, prompt_embeddings, prompt_embeddings)
            pooled = out.mean(dim=1)
            _ = classification_head(pooled)
        if device.type == "cuda":
            torch.cuda.synchronize()
        end = time.time()

    latency_ms = (end - start) * 1000.0 / iters
    peak_mem_gb = None
    if device.type == "cuda":
        peak_mem_gb = torch.cuda.max_memory_allocated() / (1024**3)

    return {
        "total_params": total_params,
        "trainable_params": trainable_params,
        "latency_ms_per_forward": latency_ms,
        "peak_mem_gb": peak_mem_gb,
    }


def _cache_path_for_file(file_path, prompt_mode):
    cache_dir = Path(__file__).resolve().parents[2] / "prompt_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(file_path).stem
    return cache_dir / f"{stem}__{prompt_mode}__last_hidden_state_fp16.pt"


def build_or_load_prompt_cache(file_path, all_data, device, prompt_mode="original", max_length=50, batch_size=64):
    """
    Build (or load) cached BERT last_hidden_state for ALL samples in a given chunk file.

    Saves token-level embeddings: (N, L, 768) in FP16 on CPU.
    """
    import torch

    cache_path = _cache_path_for_file(file_path, prompt_mode)
    if cache_path.exists():
        return torch.load(cache_path, map_location="cpu")

    # all_data: numpy array (N, C, T)
    X = torch.tensor(all_data, dtype=torch.float32)  # CPU

    N = X.size(0)
    min_values, max_values, median_values, trends = generate_statistics(X)
    prompts = generate_prompts(N, min_values, max_values, median_values, trends, mode=prompt_mode)

    bert_model.eval()
    all_embeds = []
    with torch.no_grad():
        start_idx = 0
        current_bs = max(1, int(batch_size))
        while start_idx < N:
            local_bs = min(current_bs, N - start_idx)
            p_batch = prompts[start_idx:start_idx + local_bs]
            try:
                prompt_inputs = tokenizer(
                    p_batch, return_tensors="pt", padding="max_length", truncation=True, max_length=max_length
                ).to(device)
                embeds = bert_model(**prompt_inputs).last_hidden_state
                all_embeds.append(embeds.detach().cpu().half())
                start_idx += local_bs
            except Exception as exc:
                if device.type == "cuda" and _is_cuda_oom_error(exc) and local_bs > 1:
                    torch.cuda.empty_cache()
                    current_bs = max(1, local_bs // 2)
                    print(f"[OOM保护] prompt cache batch size 降为 {current_bs}")
                    continue
                raise

    cache_tensor = torch.cat(all_embeds, dim=0)  # (N, L, 768), fp16, cpu
    torch.save(cache_tensor, cache_path)
    return cache_tensor


def train_model(model_components, dataloader, optimizer, criterion, device, num_labels, prompt_mode="original", soft_prompt=None, prompt_cache_cpu=None, use_cached_prompts=False):
    bert_model, patch_embedding, reprogramming_layer, classification_head = model_components
    bert_model.eval()
    patch_embedding.train()
    reprogramming_layer.train()
    classification_head.train()
    total_loss = 0

    for batch in dataloader:
        if len(batch) == 3:
            batch_eeg, batch_labels, batch_idx = batch
        else:
            batch_eeg, batch_labels = batch
            batch_idx = None
        try:
            batch_eeg = batch_eeg.to(device)
            batch_labels = batch_labels.to(device)

            optimizer.zero_grad()

            eeg_embeddings = patch_embedding(batch_eeg)

            min_values, max_values, median_values, trends = generate_statistics(batch_eeg)
            prompts = generate_prompts(batch_eeg.size(0), min_values, max_values, median_values, trends, mode=prompt_mode)
            if prompt_mode == "soft":
                if soft_prompt is None:
                    raise ValueError("soft_prompt must be provided when prompt_mode='soft'")
                prompt_embeddings = soft_prompt.expand(batch_eeg.size(0), -1, -1)  # (B, L, 768)
            else:
                if use_cached_prompts and (prompt_cache_cpu is not None) and (batch_idx is not None):
                    prompt_embeddings = prompt_cache_cpu[batch_idx].to(device).float()
                else:
                    prompt_inputs = tokenizer(
                        prompts, return_tensors="pt", padding="max_length", truncation=True, max_length=50
                    ).to(device)
                    with torch.no_grad():
                        prompt_embeddings = bert_model(**prompt_inputs).last_hidden_state  # (B, L, 768)
            eeg_embeddings, _ = reprogramming_layer(eeg_embeddings, prompt_embeddings, prompt_embeddings)

            pooled_output = eeg_embeddings.mean(dim=1)

            logits = classification_head(pooled_output)
            loss = criterion(logits, batch_labels)

            loss.backward()

            optimizer.step()

            total_loss += loss.item()
        except Exception as exc:
            if device.type == "cuda" and _is_cuda_oom_error(exc):
                optimizer.zero_grad(set_to_none=True)
                torch.cuda.empty_cache()
                print("[OOM保护] 训练批次发生CUDA OOM，已跳过该批次。")
                continue
            raise

    average_loss = total_loss / len(dataloader)
    return average_loss


def evaluate_model(
    model_components,
    dataloader,
    device,
    num_labels,
    prompt_mode="original",
    soft_prompt=None,
    save_attn_dir=None,
    max_attn_batches=1,
    prompt_cache_cpu=None,
    use_cached_prompts=False,
    save_attn_npz_path=None,
    return_timing=False,
):
    bert_model, patch_embedding, reprogramming_layer, classification_head = model_components
    bert_model.eval()
    patch_embedding.eval()
    reprogramming_layer.eval()
    classification_head.eval()

    all_labels = []
    all_predictions = []
    all_attn_blocks = []
    all_token_id_blocks = []
    all_label_blocks = []
    eval_batches = 0

    _sync_cuda_if_needed(device)
    infer_start = time.perf_counter()

    with torch.no_grad():
        for batch in dataloader:
            if len(batch) == 3:
                batch_eeg, batch_labels, batch_idx = batch
            else:
                batch_eeg, batch_labels = batch
                batch_idx = None
            try:
                batch_eeg = batch_eeg.to(device)
                batch_labels = batch_labels.to(device)

                eeg_embeddings = patch_embedding(batch_eeg)

                min_values, max_values, median_values, trends = generate_statistics(batch_eeg)
                prompts = generate_prompts(batch_eeg.size(0), min_values, max_values, median_values, trends, mode=prompt_mode)
                prompt_inputs = None
                if prompt_mode != "soft":
                    need_prompt_inputs = (
                        (save_attn_npz_path is not None)
                        or not (use_cached_prompts and (prompt_cache_cpu is not None) and (batch_idx is not None))
                    )
                    if need_prompt_inputs:
                        prompt_inputs = tokenizer(
                            prompts, return_tensors="pt", padding="max_length", truncation=True, max_length=50
                        )
                if prompt_mode == "soft":
                    if soft_prompt is None:
                        raise ValueError("soft_prompt must be provided when prompt_mode='soft'")
                    prompt_embeddings = soft_prompt.expand(batch_eeg.size(0), -1, -1)
                else:
                    if use_cached_prompts and (prompt_cache_cpu is not None) and (batch_idx is not None):
                        prompt_embeddings = prompt_cache_cpu[batch_idx].to(device).float()
                    else:
                        if prompt_inputs is None:
                            prompt_inputs = tokenizer(
                                prompts, return_tensors="pt", padding="max_length", truncation=True, max_length=50
                            )
                        prompt_embeddings = bert_model(**prompt_inputs.to(device)).last_hidden_state
                eeg_embeddings, attn_weights = reprogramming_layer(eeg_embeddings, prompt_embeddings, prompt_embeddings)

                # Optionally save attention maps for interpretability
                if save_attn_dir is not None and max_attn_batches > 0:
                    os.makedirs(save_attn_dir, exist_ok=True)
                    # save only a few batches to avoid huge files
                    if len(os.listdir(save_attn_dir)) < max_attn_batches:
                        try:
                            save_attention_heatmap(attn_weights, prompts, os.path.join(save_attn_dir, f"attn_batch{len(os.listdir(save_attn_dir))}.png"))
                        except Exception:
                            pass

                if save_attn_npz_path is not None:
                    attn_mean_head = attn_weights.mean(dim=1).detach().cpu().numpy().astype(np.float16)
                    if prompt_mode == "soft":
                        token_ids = np.full((attn_mean_head.shape[0], attn_mean_head.shape[-1]), -1, dtype=np.int32)
                    else:
                        if prompt_inputs is None:
                            prompt_inputs = tokenizer(
                                prompts, return_tensors="pt", padding="max_length", truncation=True, max_length=50
                            )
                        token_ids = prompt_inputs["input_ids"].cpu().numpy().astype(np.int32)
                    all_attn_blocks.append(attn_mean_head)
                    all_token_id_blocks.append(token_ids)
                    all_label_blocks.append(batch_labels.detach().cpu().numpy().astype(np.int64))

                pooled_output = eeg_embeddings.mean(dim=1)

                logits = classification_head(pooled_output)
                predictions = torch.argmax(logits, dim=1)

                all_labels.extend(batch_labels.cpu().numpy())
                all_predictions.extend(predictions.cpu().numpy())
                eval_batches += 1
            except Exception as exc:
                if device.type == "cuda" and _is_cuda_oom_error(exc):
                    torch.cuda.empty_cache()
                    print("[OOM保护] 评估批次发生CUDA OOM，已跳过该批次。")
                    continue
                raise

    _sync_cuda_if_needed(device)
    infer_time_sec = time.perf_counter() - infer_start

    if not all_labels:
        raise RuntimeError("评估阶段全部批次因 OOM 被跳过，请减小 --batch_size。")

    save_attention_records(all_attn_blocks, all_token_id_blocks, all_label_blocks, save_attn_npz_path)

    _ = classification_report(all_labels, all_predictions, zero_division=0)

    accuracy = accuracy_score(all_labels, all_predictions)
    f1 = f1_score(all_labels, all_predictions, average='macro', zero_division=0)
    precision = precision_score(all_labels, all_predictions, average='macro', zero_division=0)
    recall = recall_score(all_labels, all_predictions, average='macro', zero_division=0)
    conf_matrix = confusion_matrix(all_labels, all_predictions)
    infer_samples = len(all_labels)
    throughput = infer_samples / infer_time_sec if infer_time_sec > 0 else 0.0

    if return_timing:
        timing_info = {
            "inference_time_sec": float(infer_time_sec),
            "inference_samples": int(infer_samples),
            "inference_batches": int(eval_batches),
            "throughput_samples_per_sec": float(throughput),
        }
        return accuracy, f1, precision, recall, conf_matrix, timing_info

    return accuracy, f1, precision, recall, conf_matrix


def check_class_distribution(labels, dataset_name="Dataset", log_file=None):
    class_counts = Counter(labels)
    if log_file:
        print(f"{dataset_name} class distribution: {class_counts}")
        log_file.write(f"{dataset_name} class distribution: {class_counts}\n")
    else:
        print(f"{dataset_name} class distribution: {class_counts}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--data_dir",
        default=os.getenv(
            "SEED_CHUNKS_DIR",
            DEFAULT_DATA_DIR,
        ),
        help="Directory containing preprocessed SEED chunk .npz files",
    )
    ap.add_argument(
        "--bert",
        default=DEFAULT_BERT_DIR,
        help="Local BERT model directory",
    )
    ap.add_argument(
        "--batch_size",
        type=int,
        default=int(os.getenv("TRAIN_BATCH_SIZE", "16")),
        help="Training/evaluation batch size (default=16, lower to reduce OOM risk).",
    )
    ap.add_argument(
        "--cache_batch_size",
        type=int,
        default=int(os.getenv("CACHE_BATCH_SIZE", "16")),
        help="Batch size used when precomputing prompt cache (default=16).",
    )
    ap.add_argument(
        "--profile_batch_size",
        type=int,
        default=int(os.getenv("PROFILE_BATCH_SIZE", "8")),
        help="Batch size for efficiency profiling.",
    )
    ap.add_argument(
        "--profile_iters",
        type=int,
        default=int(os.getenv("PROFILE_ITERS", "50")),
        help="Iterations for efficiency profiling.",
    )
    ap.add_argument(
        "--profile_warmup",
        type=int,
        default=int(os.getenv("PROFILE_WARMUP", "10")),
        help="Warmup iterations for efficiency profiling.",
    )
    args = ap.parse_args()

    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')

    log_path = Path(__file__).resolve().parents[2] / "results_confusion_matrix.xlsx"
    runtime_path = Path(__file__).resolve().parents[2] / "runtime_metrics.jsonl"
    with XlsxTextLogger(log_path) as log_file:
        global bert_model, tokenizer
        bert_path = _resolve_bert_source(args.bert)

        _validate_local_bert_weights(bert_path)
        tokenizer = BertTokenizer.from_pretrained(
            bert_path,
            local_files_only=True,
        )
        try:
            bert_model = BertModel.from_pretrained(
                bert_path,
                weights_only=False,
                local_files_only=True,
            ).to(device)
        except TypeError:
            bert_model = BertModel.from_pretrained(
                bert_path,
                local_files_only=True,
            ).to(device)

        for param in bert_model.parameters():
            param.requires_grad = False

        data_root = Path(args.data_dir).expanduser()
        if not data_root.exists() and Path("/home/xiaoying/SEED-IV_chunks").exists():
            data_root = Path("/home/xiaoying/SEED-IV_chunks")
        if not data_root.exists():
            raise FileNotFoundError(f"未找到数据目录: {data_root}")
        preprocessed_files = sorted([f for f in data_root.iterdir() if f.suffix == '.npz'])
        print(f"Found {len(preprocessed_files)} files in {data_root}")
        log_file.write(f"Found {len(preprocessed_files)} files in {data_root}\n")
        if not preprocessed_files:
            raise RuntimeError(
                f"目录 {data_root} 下未找到 .npz 文件。"
                "请通过 --data_dir 指向真实的 SEED_chunks 路径。"
            )

        metrics_per_fold = []

        for file_path in preprocessed_files:
            subject_wall_start = time.perf_counter()
            subject_train_time_sec = 0.0
            subject_epochs = 0
            subject_folds = 0
            print(f"\nProcessing file: {file_path}")
            log_file.write(f"\nProcessing file: {file_path}\n")

            eeg_dataset = EEGDataset(file_path)

            all_data = eeg_dataset.data
            all_labels = eeg_dataset.labels.astype(np.int64)
            num_labels = len(np.unique(all_labels))

            if np.any(all_labels < 0) or np.any(all_labels >= num_labels):
                print(f"Error: Labels in {file_path} are out of range [0, {num_labels - 1}]")
                continue

            num_samples = len(eeg_dataset)
            print(f"Number of samples in this file: {num_samples}")

            check_class_distribution(all_labels, f"Dataset for {file_path.name}", log_file)

            X_train_val, X_test, y_train_val, y_test = train_test_split(
                all_data, all_labels, test_size=0.2, random_state=42
            )
            all_idx = np.arange(len(all_data))
            idx_train_val, idx_test = train_test_split(
                all_idx, test_size=0.2, random_state=42
            )
            check_class_distribution(y_train_val, "Train_Val Dataset", log_file)
            check_class_distribution(y_test, "Test Dataset", log_file)

            # Prompt ablation mode: one of {"original","generic","shuffle","random","soft"}.
            prompt_mode = os.getenv("PROMPT_MODE", "original").strip().lower()
            # Prompt embedding mode:
            #   - online: run BERT forward each iteration (slow)
            #   - cached: load per-sample cached last_hidden_state (recommended)
            prompt_emb_mode = os.getenv("PROMPT_EMB_MODE", "cached").strip().lower()
            use_cached_prompts = (prompt_emb_mode == "cached")

            prompt_cache_cpu = None
            if use_cached_prompts and prompt_mode != "soft":
                prompt_cache_cpu = build_or_load_prompt_cache(
                    file_path=file_path,
                    all_data=all_data,
                    device=device,
                    prompt_mode=prompt_mode,
                    max_length=50,
                    batch_size=max(1, args.cache_batch_size),
                )

            batch_size = max(1, args.batch_size)

            test_dataset = TensorDataset(
                torch.tensor(X_test, dtype=torch.float32),
                torch.tensor(y_test, dtype=torch.long),
                torch.tensor(idx_test, dtype=torch.long),
            )
            test_loader = DataLoader(test_dataset, batch_size=batch_size, shuffle=False)

            fold = 1

            kf = KFold(n_splits=5, shuffle=True, random_state=42)
            data_indices = np.arange(len(X_train_val))
            for train_indices, val_indices in kf.split(data_indices):
                current_fold_id = int(fold)
                fold_train_time_sec = 0.0
                fold_epochs = 0
                print(f'Fold {current_fold_id}:')
                log_file.write(f'Fold {current_fold_id}:')
                X_train, X_val = X_train_val[train_indices], X_train_val[val_indices]
                idx_train = idx_train_val[train_indices]
                idx_val = idx_train_val[val_indices]
                y_train, y_val = y_train_val[train_indices], y_train_val[val_indices]

                train_dataset = TensorDataset(
                    torch.tensor(X_train, dtype=torch.float32),
                    torch.tensor(y_train, dtype=torch.long),
                    torch.tensor(idx_train, dtype=torch.long),
                )
                val_dataset = TensorDataset(
                    torch.tensor(X_val, dtype=torch.float32),
                    torch.tensor(y_val, dtype=torch.long),
                    torch.tensor(idx_val, dtype=torch.long),
                )

                train_loader = DataLoader(train_dataset, batch_size=batch_size, shuffle=True)
                val_loader = DataLoader(val_dataset, batch_size=batch_size, shuffle=False)

                class_sample_counts = np.array([len(np.where(y_train == t)[0]) for t in np.unique(y_train)])
                class_weights = 1. / torch.tensor(class_sample_counts, dtype=torch.float32).to(device)
                class_weights = class_weights / class_weights.sum() * len(np.unique(y_train))
                print(class_weights)
                log_file.write(f"{class_weights}")

                criterion = nn.CrossEntropyLoss(weight=class_weights)

                patch_embedding = PatchEmbedding(patch_len=500, d_model=128, stride=250, num_channels=62).to(device)
                reprogramming_layer = ReprogrammingLayer(embed_dim=128, llm_embed_dim=768, num_heads=8, max_len=5000).to(device)
                classification_head = ClassificationHead(llm_embed_dim=768, num_labels=len(np.unique(y_train))).to(device)

                # Learnable soft prompt (capacity-controlled baseline). Only used when prompt_mode == "soft".
                soft_prompt = None
                if prompt_mode == "soft":
                    # fixed length to match tokenizer max_length used elsewhere
                    soft_prompt = nn.Parameter(torch.randn(1, 50, 768, device=device) * 0.02)


                model_components = (bert_model, patch_embedding, reprogramming_layer, classification_head)
                # ---- Efficiency logging -> efficiency.jsonl (one line per fold) ----
                try:
                    eff_path = Path(__file__).resolve().parents[2] / "efficiency.jsonl"
                    eff = profile_efficiency(
                        model_components, device,
                        batch_size=max(1, args.profile_batch_size),
                        iters=max(1, args.profile_iters),
                        warmup=max(0, args.profile_warmup),
                        prompt_mode=prompt_mode, soft_prompt=soft_prompt,
                        prompt_cache_cpu=prompt_cache_cpu, use_cached_prompts=use_cached_prompts,
                    )
                    record = {
                        "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                        "dataset_file": str(file_path),
                        "prompt_mode": prompt_mode,
                        "prompt_emb_mode": prompt_emb_mode,
                        "fold_id": int(fold),
                        "num_labels": int(len(np.unique(y_train))),
                        **eff,
                    }
                    with open(eff_path, "a", encoding="utf-8") as f_eff:
                        f_eff.write(json.dumps(record, ensure_ascii=False) + "\n")
                except Exception:
                    pass


                optimizer = AdamW([
                    {'params': classification_head.parameters(), 'lr': 1e-5},
                    {'params': patch_embedding.parameters(), 'lr': 1e-5},
                    {'params': reprogramming_layer.parameters(), 'lr': 1e-5}
                ], weight_decay=0.01)

                best_val_accuracy = 0.0
                best_model_state = None
                no_improve_epochs = 0

                patience = 5
                num_epochs = 100

                for epoch in range(num_epochs):
                    _sync_cuda_if_needed(device)
                    epoch_train_start = time.perf_counter()
                    train_loss = train_model(
                        model_components,
                        train_loader, optimizer, criterion, device, len(np.unique(y_train)),
                        prompt_mode=prompt_mode, soft_prompt=soft_prompt, prompt_cache_cpu=prompt_cache_cpu, use_cached_prompts=use_cached_prompts,
                    )
                    _sync_cuda_if_needed(device)
                    epoch_train_time_sec = time.perf_counter() - epoch_train_start
                    fold_train_time_sec += epoch_train_time_sec
                    subject_train_time_sec += epoch_train_time_sec
                    fold_epochs += 1
                    subject_epochs += 1

                    _append_jsonl(
                        runtime_path,
                        {
                            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                            "scope": "epoch_train",
                            "dataset_file": str(file_path),
                            "prompt_mode": prompt_mode,
                            "prompt_emb_mode": prompt_emb_mode,
                            "fold_id": current_fold_id,
                            "epoch_id": int(epoch + 1),
                            "train_samples": int(len(train_dataset)),
                            "train_batches": int(len(train_loader)),
                            "epoch_train_time_sec": float(epoch_train_time_sec),
                        },
                    )

                    val_accuracy, val_f1, val_precision, val_recall, val_conf_matrix = evaluate_model(
                        model_components,
                        val_loader, device, len(np.unique(y_train)),
                        prompt_mode=prompt_mode, soft_prompt=soft_prompt, prompt_cache_cpu=prompt_cache_cpu, use_cached_prompts=use_cached_prompts,
                    )
                    print(
                        f"    Epoch {epoch + 1}/{num_epochs}, Train Loss: {train_loss}, TrainTime(s): {epoch_train_time_sec:.3f}, "
                        f"Validation Accuracy: {val_accuracy}, F1: {val_f1}, "
                        f"Precision: {val_precision}, Recall: {val_recall}, Conf_matrix: {val_conf_matrix}"
                    )
                    log_file.write(
                        f"    Epoch {epoch + 1}/{num_epochs}, Train Loss: {train_loss}, TrainTime(s): {epoch_train_time_sec:.3f}, "
                        f"Validation Accuracy: {val_accuracy}, F1: {val_f1}, "
                        f"Precision: {val_precision}, Recall: {val_recall}, Conf_matrix: {val_conf_matrix}\n"
                    )

                    if val_accuracy >= best_val_accuracy:
                        best_val_accuracy = val_accuracy
                        best_model_state = {
                            'patch_embedding': copy.deepcopy(patch_embedding.state_dict()),
                            'reprogramming_layer': copy.deepcopy(reprogramming_layer.state_dict()),
                            'classification_head': copy.deepcopy(classification_head.state_dict()),
                        }
                        no_improve_epochs = 0
                    else:
                        no_improve_epochs += 1
                        if no_improve_epochs >= patience:
                            print(f"    Early stopping on epoch {epoch + 1}")
                            log_file.write(f"    Early stopping on epoch {epoch + 1}\n")
                            break

                if best_model_state is not None:
                    patch_embedding.load_state_dict(best_model_state['patch_embedding'])
                    reprogramming_layer.load_state_dict(best_model_state['reprogramming_layer'])
                    classification_head.load_state_dict(best_model_state['classification_head'])

                    val_accuracy, val_f1, val_precision, val_recall, val_conf_matrix = evaluate_model(
                        model_components,
                        val_loader, device, len(np.unique(y_train)),
                        prompt_mode=prompt_mode, soft_prompt=soft_prompt, prompt_cache_cpu=prompt_cache_cpu, use_cached_prompts=use_cached_prompts,
                    )
                    print(
                        f"    Best Validation Accuracy: {val_accuracy}, F1: {val_f1}, "
                        f"Precision: {val_precision}, Recall: {val_recall}, Conf_matrix: {val_conf_matrix}"
                    )
                    log_file.write(
                        f"    Best Validation Accuracy: {val_accuracy}, F1: {val_f1}, "
                        f"Precision: {val_precision}, Recall: {val_recall}, Conf_matrix: {val_conf_matrix}\n"
                    )
                    _append_jsonl(
                        runtime_path,
                        {
                            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                            "scope": "fold_train",
                            "dataset_file": str(file_path),
                            "prompt_mode": prompt_mode,
                            "prompt_emb_mode": prompt_emb_mode,
                            "fold_id": current_fold_id,
                            "epochs_completed": int(fold_epochs),
                            "train_samples": int(len(train_dataset)),
                            "train_batches": int(len(train_loader)),
                            "fold_train_time_sec": float(fold_train_time_sec),
                        },
                    )
                    print(f"    Fold {current_fold_id} TrainTime(s): {fold_train_time_sec:.3f}, Epochs: {fold_epochs}")
                    log_file.write(f"    Fold {current_fold_id} TrainTime(s): {fold_train_time_sec:.3f}, Epochs: {fold_epochs}\n")
                    subject_folds += 1

                    fold += 1

            attn_record_path = (
                Path(__file__).resolve().parents[2]
                / "attn_records"
                / prompt_mode
                / f"{Path(file_path).stem}__test_attn.npz"
            )
            test_accuracy, test_f1, test_precision, test_recall, test_conf_matrix, test_timing = evaluate_model(
                model_components,
                test_loader,
                device,
                len(np.unique(y_train)),
                prompt_mode=prompt_mode,
                soft_prompt=soft_prompt,
                prompt_cache_cpu=prompt_cache_cpu,
                use_cached_prompts=use_cached_prompts,
                save_attn_dir=str(Path(__file__).resolve().parents[2] / "attn_viz"),
                max_attn_batches=1,
                save_attn_npz_path=str(attn_record_path),
                return_timing=True,
            )
            print(
                f"*** Test Dataset Results - Accuracy: {test_accuracy}, F1: {test_f1}, "
                f"Precision: {test_precision}, Recall: {test_recall}, Conf_matrix: {test_conf_matrix}, "
                f"InferTime(s): {test_timing['inference_time_sec']:.3f}, Throughput(samples/s): {test_timing['throughput_samples_per_sec']:.3f}"
            )
            log_file.write(
                f"*** Test Dataset Results - Accuracy: {test_accuracy}, F1: {test_f1}, "
                f"Precision: {test_precision}, Recall: {test_recall}, Conf_matrix: {test_conf_matrix}, "
                f"InferTime(s): {test_timing['inference_time_sec']:.3f}, Throughput(samples/s): {test_timing['throughput_samples_per_sec']:.3f}\n"
            )
            _append_jsonl(
                runtime_path,
                {
                    "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                    "scope": "test_inference",
                    "dataset_file": str(file_path),
                    "prompt_mode": prompt_mode,
                    "prompt_emb_mode": prompt_emb_mode,
                    "inference_time_sec": float(test_timing["inference_time_sec"]),
                    "inference_samples": int(test_timing["inference_samples"]),
                    "inference_batches": int(test_timing["inference_batches"]),
                    "throughput_samples_per_sec": float(test_timing["throughput_samples_per_sec"]),
                },
            )

            subject_wall_time_sec = time.perf_counter() - subject_wall_start
            _append_jsonl(
                runtime_path,
                {
                    "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
                    "scope": "subject_train_summary",
                    "dataset_file": str(file_path),
                    "prompt_mode": prompt_mode,
                    "prompt_emb_mode": prompt_emb_mode,
                    "subject_folds": int(subject_folds),
                    "subject_epochs": int(subject_epochs),
                    "subject_train_time_sec": float(subject_train_time_sec),
                    "subject_total_wall_time_sec": float(subject_wall_time_sec),
                },
            )
            print(
                f"*** Subject Runtime - TrainTime(s): {subject_train_time_sec:.3f}, "
                f"TotalWallTime(s): {subject_wall_time_sec:.3f}, Folds: {subject_folds}, Epochs: {subject_epochs}"
            )
            log_file.write(
                f"*** Subject Runtime - TrainTime(s): {subject_train_time_sec:.3f}, "
                f"TotalWallTime(s): {subject_wall_time_sec:.3f}, Folds: {subject_folds}, Epochs: {subject_epochs}\n"
            )

            metrics_per_fold.append({
                'accuracy': test_accuracy,
                'f1': test_f1,
                'precision': test_precision,
                'recall': test_recall,
                'conf_matrix': test_conf_matrix
            })

        if metrics_per_fold:
            avg_accuracy = np.mean([m['accuracy'] for m in metrics_per_fold])
            avg_f1 = np.mean([m['f1'] for m in metrics_per_fold])
            avg_precision = np.mean([m['precision'] for m in metrics_per_fold])
            avg_recall = np.mean([m['recall'] for m in metrics_per_fold])
            avg_conf_matrix = np.mean([m['conf_matrix'] for m in metrics_per_fold])

            print("Average Metrics across all files:")
            print(f"  Accuracy: {avg_accuracy:.2f}")
            print(f"  F1 Score: {avg_f1:.2f}")
            print(f"  Precision: {avg_precision:.2f}")
            print(f"  Recall: {avg_recall:.2f}")
            print(f"  Conf_matrix: {avg_conf_matrix:.2f}")

            log_file.write("Average Metrics across all files:\n")
            log_file.write(f"  Accuracy: {avg_accuracy:.2f}\n")
            log_file.write(f"  F1 Score: {avg_f1:.2f}\n")
            log_file.write(f"  Precision: {avg_precision:.2f}\n")
            log_file.write(f"  Recall: {avg_recall:.2f}\n")
            log_file.write(f"  Conf_matrix: {avg_conf_matrix:.2f}")


if __name__ == '__main__':
    main()
